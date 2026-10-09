"""Format-neutral CSV/XLSX parsing, mapping, normalization and validation."""
import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from typing import Any
from openpyxl import load_workbook
from app.taxation.domain import TAX_CATEGORIES

MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10000
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
ALLOWED_MIME_TYPES = {".csv": {"text/csv", "application/csv", "text/plain", None}, ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None}}
DECIMAL_FIELDS = {"subtotal", "taxable_base", "total_tax", "total", "rate", "amount"}
STRING_FIELDS = {"document_type", "direction", "series", "document_number", "authorization_number", "currency", "status", "jurisdiction"}
TAX_FIELDS = "tax_code|tax_name|tax_type|category|rate|taxable_base|amount|is_withholding|retention_id|withholding_type|withholding_role|applied_to|withholding_effect"
PATH_RE = re.compile(rf"^taxes\[(\d+)\]\.({TAX_FIELDS})$")
LINE_RE = re.compile(rf"^lines\[(\d+)\]\.((?:taxes\[(\d+)\]\.)?[^.]+)$")
DATE_FORMATS = {"DMY": "%d/%m/%Y", "MDY": "%m/%d/%Y", "YMD": "%Y/%m/%d"}

class ImportFormatError(ValueError):
    pass

def validate_upload(filename: str | None, content: bytes, content_type: str | None = None) -> str:
    if not filename or any(ord(char) < 32 for char in filename) or len(PurePath(filename).name) > 255:
        raise ImportFormatError("Invalid filename")

    suffix = PurePath(filename).suffix.lower()

    if suffix not in ALLOWED_EXTENSIONS: raise ImportFormatError("Unsupported file extension")

    if len(content) > MAX_IMPORT_FILE_BYTES: raise ImportFormatError("File exceeds the maximum import size")

    if content_type not in ALLOWED_MIME_TYPES[suffix]: raise ImportFormatError("File MIME type does not match the supported format")

    if suffix == ".xlsx":
        if not zipfile.is_zipfile(io.BytesIO(content)): raise ImportFormatError("Invalid XLSX file")

    return "XLSX" if suffix == ".xlsx" else "CSV"

def read_rows(content: bytes, file_format: str, *, delimiter: str | None = None, sheet: str | None = None, encoding: str | None = None, header_row: int = 1, data_start_row: int | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    if header_row < 1 or (data_start_row is not None and data_start_row < header_row):
        raise ImportFormatError("Invalid header/data row configuration")

    if file_format == "CSV":
        try:
            text = content.decode(encoding or "utf-8-sig")

        except (LookupError, UnicodeDecodeError) as exc:
            raise ImportFormatError("Unsupported or invalid file encoding") from exc

        if not text.strip(): raise ImportFormatError("CSV is empty")

        lines = text.splitlines()

        if len(lines) < header_row:
            raise ImportFormatError("CSV header row was not found")

        selected = lines[header_row - 1:]

        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff("\n".join(selected[:20]), delimiters=",;\t|").delimiter

            except csv.Error as exc:
                raise ImportFormatError("Unable to detect CSV delimiter") from exc

        reader = csv.reader(selected, delimiter=delimiter)
        headers = [header.strip() if header else "" for header in next(reader, [])]
        skip = (data_start_row - header_row - 1) if data_start_row else 0

        for _ in range(skip): next(reader, None)

        rows = [dict(zip(headers, row)) for row in reader]

    elif file_format == "XLSX":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet_name = sheet or workbook.sheetnames[0]

            if sheet_name not in workbook.sheetnames: raise ImportFormatError("Requested worksheet was not found")

            if any(isinstance(cell.value, str) and cell.value.startswith("=") for row in formulas[sheet_name].iter_rows() for cell in row):
                raise ImportFormatError("Workbook formulas are not accepted")

            values = workbook[sheet_name].iter_rows(values_only=True)

            for _ in range(header_row - 1): next(values, ())

            headers = [str(value).strip() if value is not None else "" for value in next(values, ())]
            skip = (data_start_row - header_row - 1) if data_start_row else 0

            for _ in range(skip): next(values, ())

            rows = [dict(zip(headers, row)) for row in values]

        except ImportFormatError: raise
        except Exception as exc: raise ImportFormatError("Invalid XLSX content") from exc

    else: raise ImportFormatError("Unsupported detected format")

    if not headers or any(not header for header in headers) or len(headers) != len(set(headers)): raise ImportFormatError("Headers must be present and unique")

    if not rows: raise ImportFormatError("File has no data rows")

    if len(rows) > MAX_IMPORT_ROWS: raise ImportFormatError("File exceeds the maximum number of rows")

    return headers, rows

def suggest_mapping(headers: list[str]) -> dict[str, str]:
    known = {"document_type", "direction", "issue_date", "document_date", "series", "document_number", "number", "authorization_number", "authorization", "currency", "subtotal", "gross_amount", "taxable_base", "total_tax", "tax_amount", "total", "total_amount", "status", "jurisdiction", "issuer.name", "issuer.tax_identifier", "issuer.tax_identifier_type", "receiver.name", "receiver.tax_identifier", "receiver.tax_identifier_type"}
    normalized = {re.sub(r"[^a-z0-9_.]+", "_", header.strip().lower()).strip("_"): header for header in headers}
    result = {}

    for target in known:
        source = normalized.get(target.replace(".", "_"))

        if source: result[source] = target

    return result

def normalize_and_validate(rows: list[dict[str, Any]], mapping: dict[str, str], *, locale: str | None = None, date_format: str | None = None, decimal_separator: str | None = None, thousands_separator: str | None = None, group_by: list[str] | None = None) -> dict[str, Any]:
    normalized_rows, errors, warnings = [], [], []

    for row_number, row in enumerate(rows, start=2):
        try:
            normalized = _map_row(row, mapping, locale, date_format, decimal_separator, thousands_separator)
            row_errors, row_warnings = _validate_row(normalized)

        except ImportFormatError:
            normalized, row_errors, row_warnings = {}, ["Invalid row data"], []

        if row_errors: errors.append({"row": row_number, "errors": row_errors})

        if row_warnings: warnings.append({"row": row_number, "warnings": row_warnings})

        normalized_rows.append({"row": row_number, "data": normalized, "original": dict(row), "errors": row_errors, "warnings": row_warnings})

    if group_by:
        normalized_rows = _group_rows(normalized_rows, group_by)
        errors = [error for item in normalized_rows for error in item["errors"]]
        warnings = [warning for item in normalized_rows for warning in item["warnings"]]

    return {"rows": normalized_rows, "errors": errors, "warnings": warnings, "total": len(rows),
            "rows_detected": len(rows), "rows_valid": sum(not item["errors"] for item in normalized_rows),
            "rows_invalid": sum(bool(item["errors"]) for item in normalized_rows)}

def _map_row(row, mapping, locale, date_format=None, decimal_separator=None, thousands_separator=None):
    result: dict[str, Any] = {"taxes": []}

    for source, target in mapping.items():
        if source not in row: continue

        value = _clean_value(row[source])

        if value is None: continue

        target = _canonical_target(target)
        match = PATH_RE.match(target)

        if match:
            index, field = int(match.group(1)), match.group(2)

            while len(result["taxes"]) <= index: result["taxes"].append({})

            result["taxes"][index][field] = _normalize_field(field, value, locale, date_format, decimal_separator, thousands_separator)

        elif (line_match := LINE_RE.match(target)):
            line_index = int(line_match.group(1)); nested = line_match.group(2)
            result.setdefault("lines", [])

            while len(result["lines"]) <= line_index: result["lines"].append({"taxes": []})

            tax_match = re.match(rf"^taxes\[(\d+)\]\.({TAX_FIELDS})$", nested)

            if tax_match:
                tax_index, field = int(tax_match.group(1)), tax_match.group(2)

                while len(result["lines"][line_index]["taxes"]) <= tax_index: result["lines"][line_index]["taxes"].append({})

                result["lines"][line_index]["taxes"][tax_index][field] = _normalize_field(field, value, locale, date_format, decimal_separator, thousands_separator)

            else:
                result["lines"][line_index][nested] = _normalize_field(nested, value, locale, date_format, decimal_separator, thousands_separator)

        elif "." in target:
            parent, field = target.split(".", 1)
            result.setdefault(parent, {})[field] = _normalize_field(field, value, locale, date_format, decimal_separator, thousands_separator)

        else: result[target] = _normalize_field(target, value, locale, date_format, decimal_separator, thousands_separator)

    return result

def _normalize_field(field, value, locale, date_format=None, decimal_separator=None, thousands_separator=None):
    if field in DECIMAL_FIELDS: return _parse_decimal(value, decimal_separator, thousands_separator)

    if field in {"issue_date", "document_date"}: return _parse_date(value, date_format or locale)

    if field == "is_withholding":
        normalized = str(value).strip().lower()

        if normalized in {"true", "1", "yes", "si"}: return True

        if normalized in {"false", "0", "no"}: return False

        raise ImportFormatError("Invalid withholding boolean")

    if field in STRING_FIELDS or field in {"name", "tax_identifier", "tax_identifier_type", "tax_code", "tax_name", "tax_type", "category", "retention_id", "withholding_type", "withholding_role", "applied_to", "withholding_effect"}: return str(value).strip()

    return value

def _parse_date(value, locale):
    if isinstance(value, datetime): return value.date()

    if isinstance(value, date): return value

    text = str(value).strip()

    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try: return datetime.strptime(text, fmt).date()
        except ValueError: pass

    if "/" in text or "-" in text:
        separator = "/" if "/" in text else "-"
        parts = text.split(separator)

        if len(parts) == 3 and len(parts[0]) <= 2 and len(parts[1]) <= 2:

            if locale == "YMD":
                try: return datetime.strptime(text, "%Y/%m/%d" if separator == "/" else "%Y-%m-%d").date()
                except ValueError as exc: raise ImportFormatError("Invalid issue_date") from exc

            if locale not in {"DMY", "MDY"}: raise ImportFormatError("Ambiguous date requires locale DMY or MDY")

            try:
                return datetime.strptime(text, "%d/%m/%Y" if locale == "DMY" else "%m/%d/%Y").date()

            except ValueError as exc:
                raise ImportFormatError("Invalid issue_date") from exc

    raise ImportFormatError("Invalid issue_date")

def _parse_decimal(value, decimal_separator=None, thousands_separator=None) -> Decimal:
    if isinstance(value, Decimal): return value

    text = str(value).strip().replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")

    if negative: text = "-" + text[1:-1]

    text = re.sub(r"[^0-9,\.\-+]", "", text)

    if decimal_separator:
        if thousands_separator:
            text = text.replace(thousands_separator, "")

        else:
            other_separator = "." if decimal_separator == "," else ","

            if other_separator in text and decimal_separator in text:
                text = text.replace(other_separator, "")

        if decimal_separator != ".": text = text.replace(decimal_separator, ".")

    if not text: raise ImportFormatError("Empty numeric value")

    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")

    elif "," in text: text = text.replace(",", ".")

    try: result = Decimal(text)

    except InvalidOperation as exc: raise ImportFormatError("Invalid numeric value") from exc

    return result

def _canonical_target(target):
    aliases = {
        "document_date": "issue_date", "number": "document_number", "authorization": "authorization_number",
        "gross_amount": "subtotal", "total_amount": "total", "tax_amount": "total_tax",
        "tax_rate": "rate", "tax_category": "category", "tax_base": "taxable_base", "withholding_base": "taxable_base",
        "withholding_amount": "amount", "line_base": "taxable_base", "line_total": "total",
    }

    if target.startswith("taxes["):
        tax_aliases = {"tax_rate": "rate", "tax_category": "category", "tax_base": "taxable_base", "tax_amount": "amount", "withholding_base": "taxable_base", "withholding_amount": "amount"}
        
        return re.sub(r"\.(tax_rate|tax_category|tax_base|tax_amount|withholding_base|withholding_amount)$", lambda match: "." + tax_aliases[match.group(1)], target)

    if target.startswith("lines["):
        line_aliases = {"line_base": "taxable_base", "line_total": "total"}
        tax_aliases = {"tax_rate": "rate", "tax_category": "category", "tax_base": "taxable_base", "tax_amount": "amount", "withholding_base": "taxable_base", "withholding_amount": "amount"}
        target = re.sub(r"\.(tax_rate|tax_category|tax_base|tax_amount|withholding_base|withholding_amount)$", lambda match: "." + tax_aliases[match.group(1)], target)
        
        return re.sub(r"\.(line_base|line_total)$", lambda match: "." + line_aliases[match.group(1)], target)
    
    return aliases.get(target, target)

def _group_rows(rows, group_by):
    groups = {}

    for item in rows:
        key = tuple(_value_at(item["data"], _canonical_target(field)) for field in group_by)

        if key not in groups:
            groups[key] = dict(item)
            groups[key]["source_rows"] = [item["row"]]

            continue

        current = groups[key]
        current["source_rows"].append(item["row"])
        current["original"].update(item["original"])
        _merge_document_rows(current["data"], item["data"])
        current["errors"].extend(item["errors"])
        current["warnings"].extend(item["warnings"])

    return list(groups.values())

def _merge_document_rows(target, source):
    for key, value in source.items():
        if key == "lines":
            target.setdefault("lines", []).extend(value)

        elif key == "taxes":
            target.setdefault("taxes", []).extend(value)

        elif target.get(key) in (None, ""):
            target[key] = value

    if target.get("lines"):
        target.setdefault("taxable_base", sum((_decimal_or_zero(line.get("taxable_base")) for line in target["lines"]), Decimal("0")))
        target.setdefault("total", sum((_decimal_or_zero(line.get("total")) for line in target["lines"]), Decimal("0")))

def _value_at(data, path):
    value = data

    for part in path.split("."):
        value = value.get(part) if isinstance(value, dict) else None

    return value

def _decimal_or_zero(value):
    return value if isinstance(value, Decimal) else Decimal(str(value or "0"))

def _clean_value(value):
    if value is None: return None

    if isinstance(value, str):
        value = value.strip()

        return value or None

    return value

def _validate_row(data):
    errors, warnings = [], []

    for field in ("document_type", "issue_date", "currency"):
        if not data.get(field): errors.append(f"Missing required mapped field: {field}")

    if not data.get("issuer", {}).get("tax_identifier"): warnings.append("issuer.tax_identifier is empty")

    for tax in data.get("taxes", []):
        if tax.get("category") not in {None, *TAX_CATEGORIES}: errors.append("Unknown tax category")

        for field in ("rate", "taxable_base", "amount"):
            if isinstance(tax.get(field), Decimal) and tax[field] < 0:
                errors.append(f"Negative tax {field}")

    for field in ("subtotal", "taxable_base", "total_tax", "total"):
        if isinstance(data.get(field), Decimal) and data[field] < 0:
            errors.append(f"Negative monetary field: {field}")

    if data.get("subtotal") is not None and data.get("total_tax") is not None and data.get("total") is not None and data["subtotal"] + data["total_tax"] != data["total"]:
        warnings.append("subtotal plus total_tax does not equal total")

    return errors, warnings